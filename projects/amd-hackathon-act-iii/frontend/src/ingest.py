"""ingest.py: turns uploaded documents (Word, PDF, text, CSV, JSON) into the backend's two input files.
Input: file name and bytes. Output: lists in the data-contract format, ready for backend/data/sop.json and notes.json.
No Streamlit in here, so it can be tested alone. The checks mirror backend/src/stages/loader.py."""
import csv, io, json, re
from datetime import date, datetime

STEP_FIELDS = ["sop_id", "step_id", "title", "text"]
NOTE_FIELDS = ["note_id", "date", "shift", "source", "text"]
SOURCES = ["handover", "near-miss", "maintenance"]
SHIFTS = ["Day", "Evening", "Night"]
ALIASES = {"note_id": ("note_id", "id", "note id", "noteid"), "date": ("date", "day", "timestamp", "logged", "created"),
           "shift": ("shift", "crew"), "source": ("source", "type", "category", "kind"),
           "text": ("text", "note", "notes", "description", "comment", "comments", "body", "details")}
STEP_RE = re.compile(r"^\s*(?:step\s*)?(\d{1,3})\s*[.):\-]\s*(\S.*)$", re.I)
DATE_RE = re.compile(r"^\s*(\d{4}-\d{2}-\d{2}|\d{1,2}[/.]\d{1,2}[/.]\d{4})\b[\s,:;\-]*(.*)$", re.S)

def _ext(name):
    return name.lower().rsplit(".", 1)[-1]

def read_text(name, data):
    """Plain text out of a PDF, Word, text or markdown file."""
    ext = _ext(name)
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # a table row like "3 | Isolate power" becomes "3. Isolate power"
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(f"{cells[0]}. {' '.join(cells[1:])}" if cells[0].isdigit() else " | ".join(cells))
        return "\n".join(lines)
    return data.decode("utf-8-sig", errors="replace")

def read_rows(name, data):
    """Rows (dicts) out of a CSV or JSON file."""
    s = data.decode("utf-8-sig", errors="replace")
    if _ext(name) == "json":
        obj = json.loads(s)
        rows = obj if isinstance(obj, list) else next((v for v in obj.values() if isinstance(v, list)), [])
    else:
        try:
            dialect = csv.Sniffer().sniff(s[:2000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(io.StringIO(s), dialect=dialect))
    return [r for r in rows if isinstance(r, dict)]

def make_title(text):
    """Short title for a step: its first sentence or clause, at most 8 words."""
    head = re.split(r"(?<=[.;])\s|:\s", text, maxsplit=1)[0].strip(" .:;")
    words = head.split()
    return " ".join(words[:8]) + ("..." if len(words) > 8 else "")

def split_sop_steps(text):
    """Numbered lines ('1.', '2)', 'Step 3:') become steps. No numbering: one step per paragraph."""
    steps = []
    for line in text.splitlines():
        line = line.strip()
        m = STEP_RE.match(line)
        if m:
            steps.append(m.group(2))
        elif line and steps:
            steps[-1] += " " + line  # a wrapped line belongs to the step above
    if not steps:
        steps = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    steps = [" ".join(s.split()) for s in steps]
    return [(make_title(s), s) for s in steps]

def parse_sop_file(name, data, sop_id):
    """One SOP file -> step dicts in the sop.json format. CSV/JSON must already have the four columns."""
    if _ext(name) in ("csv", "json"):
        rows = read_rows(name, data)
        if rows and not all(f in {str(k).lower() for k in rows[0]} for f in STEP_FIELDS):
            raise ValueError("an SOP table needs the columns sop_id, step_id, title, text")
        return [{f: {str(k).lower(): v for k, v in r.items()}.get(f, "") for f in STEP_FIELDS} for r in rows]
    parts = split_sop_steps(read_text(name, data))
    return [{"sop_id": sop_id, "step_id": i, "title": t, "text": x} for i, (t, x) in enumerate(parts, 1)]

def to_iso(s):
    """'2026-09-01' or '01/09/2026' (day first) -> '2026-09-01'. Empty string if it is not a date."""
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return ""

def norm_source(s):
    t = re.sub(r"[^a-z]", "", str(s).lower())
    return next((k for k in SOURCES if re.sub(r"[^a-z]", "", k) in t), "")

def split_notes_text(text):
    """Free text -> raw notes. One note per paragraph, or per line that starts with a date."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if sum(bool(DATE_RE.match(l)) for l in lines) >= 2:
        blocks, cur = [], []
        for l in lines:
            if DATE_RE.match(l) and cur:
                blocks.append(" ".join(cur)); cur = []
            cur.append(l)
        blocks.append(" ".join(cur))
    else:
        blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
    out = []
    for b in blocks:
        b = " ".join(b.split()); m = DATE_RE.match(b)
        note = {"date": to_iso(m.group(1)) if m else "", "shift": "", "source": "", "text": m.group(2) if m else b}
        s = re.match(r"^(day|evening|night)\b[\s,:;\-]*(.*)$", note["text"], re.I)  # '2026-09-01 Night: ...'
        if s:
            note["shift"], note["text"] = s.group(1).title(), s.group(2)
        s = re.match(r"^(handover|near[- ]?miss|maintenance)\b[\s,:;\-]*(.*)$", note["text"], re.I)
        if s:
            note["source"], note["text"] = norm_source(s.group(1)), s.group(2)
        out.append(note)
    return out

def parse_notes_file(name, data):
    """One notes file -> raw notes (fields may be blank; finish_notes fills them)."""
    if _ext(name) in ("csv", "json"):
        out = []
        for r in read_rows(name, data):
            low = {str(k).strip().lower(): v for k, v in r.items()}
            out.append({f: str(next((low[a] for a in ALIASES[f] if low.get(a) not in (None, "")), "")).strip() for f in NOTE_FIELDS})
        return out
    return split_notes_text(read_text(name, data))

def finish_notes(raw, defaults):
    """Fill gaps with the defaults, number the notes N001, N002... and drop empty ones."""
    out = []
    for r in raw:
        text = " ".join(str(r.get("text", "")).split())
        if text:
            out.append({"note_id": str(r.get("note_id") or f"N{len(out) + 1:03d}"), "date": to_iso(r.get("date", "")) or defaults["date"],
                        "shift": str(r.get("shift", "")).strip().title() or defaults["shift"],
                        "source": norm_source(r.get("source", "")) or defaults["source"], "text": text})
    return out

def _clean(v):
    return "" if v is None or (isinstance(v, float) and v != v) else v

def clean_steps(rows):
    """Table edits -> proper step dicts (blank rows dropped, step_id whole number where possible)."""
    out = []
    for r in rows:
        r = {k: _clean(r.get(k)) for k in STEP_FIELDS}
        if any(str(v).strip() for v in r.values()):
            try:
                r["step_id"] = int(float(r["step_id"]))
            except (TypeError, ValueError):
                pass
            out.append({**r, "sop_id": str(r["sop_id"]).strip(), "title": str(r["title"]).strip(), "text": " ".join(str(r["text"]).split())})
    return out

def clean_notes(rows):
    out = []
    for r in rows:
        r = {k: str(_clean(r.get(k))).strip() for k in NOTE_FIELDS}
        if any(r.values()):
            out.append(r)
    return out

def validate(steps, notes):
    """(errors, warnings). Errors block the export. Same required fields as backend/src/stages/loader.py."""
    errors, warnings, seen, ids = [], [], set(), set()
    if not steps: errors.append("No SOP steps yet.")
    if not notes: errors.append("No notes yet.")
    for i, s in enumerate(steps, 1):
        miss = [f for f in STEP_FIELDS if str(s.get(f, "")).strip() == ""]
        if miss: errors.append(f"SOP row {i}: missing {', '.join(miss)}.")
        if not isinstance(s.get("step_id"), int): errors.append(f"SOP row {i}: step_id must be a whole number.")
        key = (s.get("sop_id"), s.get("step_id"))
        if key in seen: errors.append(f"SOP row {i}: {key[0]} step {key[1]} appears twice.")
        seen.add(key)
    for i, n in enumerate(notes, 1):
        miss = [f for f in NOTE_FIELDS if not str(n.get(f, "")).strip()]
        if miss: errors.append(f"Note row {i}: missing {', '.join(miss)}.")
        if n.get("note_id") in ids: errors.append(f"Note row {i}: note_id {n.get('note_id')} appears twice.")
        ids.add(n.get("note_id"))
        if n.get("date") and not to_iso(n["date"]) == n["date"]: errors.append(f"Note row {i}: date must look like 2026-09-01.")
        if n.get("source") and n["source"] not in SOURCES: errors.append(f"Note row {i}: type must be handover, near-miss or maintenance.")
        if n.get("shift") and n["shift"] not in SHIFTS: warnings.append(f"Note row {i}: shift '{n['shift']}' is not Day, Evening or Night.")
    return errors, warnings[:20]

def to_json(items):
    """Same style the backend writes: 2-space indent, real characters."""
    return json.dumps(items, indent=2, ensure_ascii=False)
